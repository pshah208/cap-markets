"""Coverage workflow: market shock -> triage -> outreach -> client workspace -> guardrails
-> call assist -> wrap-up.

Deterministic, from governed figures (skill_data / market_data) - the same maths as the
stress test, so the triage, the workspace and the papers can never disagree. Live agent
answers come separately from /api/chat.

The BFF runs up to three replicas, so it holds no workflow state: the browser keeps the
session's approvals and every approval is written to the log with the user's UPN.
"""

from __future__ import annotations

import io
import logging
from datetime import date
from typing import Any

from .client_skills import _band, _cushion, _ebitda_linked, _headroom, _stressed
from .market_data import MARKET_DATA
from .papers import FABRIC, FOUNDRY, WORK
from .skill_data import SKILL_DATA

logger = logging.getLogger("bff.coverage")

SCN = MARKET_DATA["scenario"]
EXTRA = MARKET_DATA["extra"]
PEOPLE = MARKET_DATA["people"]
OWNERS = MARKET_DATA["owners"]
CLIENTS = SKILL_DATA["clients"]
HERO = SKILL_DATA["client"]["id"]
AS_OF = date.fromisoformat(SCN["as_of"])
SECTOR = SKILL_DATA["sector_limit"]
SECTOR_MATCHES = lambda c: bool(SECTOR) and c["sector"].lower() in SECTOR["name"].lower()  # noqa: E731


def _days(d: str) -> int | None:
    try:
        return (date.fromisoformat(d) - AS_OF).days
    except (ValueError, TypeError):
        return None


def _sn(c: dict) -> dict | None:
    return next((l for l in c["limits"] if l["type"].startswith("Single-name")), None)


# --------------------------------------------------------------------------- #
# shock impact per client
# --------------------------------------------------------------------------- #

def impact(cid: str) -> dict[str, Any]:
    c, s = CLIENTS[cid], SCN["clients"][cid]
    shock = float(s["ebitda_shock_pct"])
    covs = []
    for cov in c["covenants"]:
        if not _ebitda_linked(cov):
            continue
        v = _stressed(cov, shock)
        h = _headroom(cov, v)
        band, tone = _band(h)
        stressed_cov = {**cov, "actual": v}
        covs.append({"covenant": cov["covenant"], "threshold": cov["threshold"], "test": cov["test"],
                     "before": cov["actual"], "after": round(v, 2), "headroom_after": round(h, 2),
                     "band_before": _band(_headroom(cov, cov["actual"]))[0], "band_after": band,
                     "tone": tone, "cushion_before": round(_cushion(cov), 1),
                     "cushion_after": round(_cushion(stressed_cov), 1) if h > 0 else 0.0})
    binding = min(covs, key=lambda x: x["cushion_after"]) if covs else None

    out = {"client_id": cid, "ebitda_shock_pct": shock, "driver": s["driver"],
           "covenants": covs, "binding": binding}
    if cid == HERO:
        # The live deal, re-cut at the shocked EBITDA.
        fin = SKILL_DATA["financials"]
        ebitda = fin["adjusted_ebitda_ltm_m"] * (1 + shock / 100)
        nd, thr = fin["total_net_debt_m"], SKILL_DATA["covenant"]["threshold"]
        rev = SKILL_DATA["opportunities"]["revolver_upsize"]
        ask = rev["proposed_commitment_total_m"] - rev["current_commitment_total_m"]
        pse = next(l for l in SKILL_DATA["limits"] if "pre-settlement" in l["type"])
        out["deal"] = {
            "ask_m": ask,
            "max_clearing_before_m": round(thr * fin["adjusted_ebitda_ltm_m"] - nd, 1),
            "max_clearing_after_m": round(max(0.0, thr * ebitda - nd), 1),
            "ask_leverage_after": round((nd + ask) / ebitda, 2),
            "resized_150_leverage_after": round((nd + 150) / ebitda, 2),
            "resized_150_clears": (nd + 150) / ebitda < thr,
        }
        after = round(pse["utilised"] + s["pse_change_m"], 1)
        out["pse"] = {"before": pse["utilised"], "after": after, "limit": pse["limit"],
                      "pct_after": round(after / pse["limit"] * 100, 1), "change": s["pse_change_m"]}
    return out


# --------------------------------------------------------------------------- #
# severity
# --------------------------------------------------------------------------- #

def _events(cid: str) -> list[str]:
    c, s = CLIENTS[cid], SCN["clients"][cid]
    ev = []
    if cid == HERO:
        m = SKILL_DATA["meeting"]
        d = _days(m["date"])
        if d is not None and 0 <= d <= 7:
            ev.append(f"meeting_7d|Client meeting {m['date']} ({d} days)")
    d = _days(c["kyc_refresh_due"])
    if d is not None and d <= 60:
        ev.append(f"kyc_60d|KYC refresh due {c['kyc_refresh_due']} ({d} days)")
    for fac in c["facilities"]:
        d = _days(fac["maturity"])
        if d is not None and 0 <= d <= 365:
            ev.append(f"maturity_12m|{fac['type']} matures {fac['maturity']}")
            break
    if s.get("pending_transaction"):
        ev.append(f"pending_transaction|{s['pending_transaction']}")
    return ev


def severity(cid: str) -> dict[str, Any]:
    w = SCN["severity"]
    c, imp = CLIENTS[cid], impact(cid)
    parts: list[dict[str, Any]] = []

    b = imp["binding"]
    if b:
        cush = b["cushion_after"]
        pts = (w["covenant"]["breach"] if b["band_after"] == "Breach" else
               w["covenant"]["lt5"] if cush < 5 else w["covenant"]["lt10"] if cush < 10 else
               w["covenant"]["lt20"] if cush < 20 else w["covenant"]["else"])
        label = ("breaches" if b["band_after"] == "Breach"
                 else f"{cush}% EBITDA cushion left")
        parts.append({"factor": "Stressed covenant", "points": pts,
                      "detail": f"{b['covenant']}: {b['before']}x -> {b['after']}x vs {b['threshold']}x - {label}"})

    sn = _sn(c)
    if sn:
        u = sn["utilisation_pct"]
        pts = w["limit"]["gt90"] if u > 90 else w["limit"]["gt80"] if u > 80 else w["limit"]["gt70"] if u > 70 else w["limit"]["else"]
        parts.append({"factor": "Single-name limit", "points": pts, "detail": f"{u}% utilised ({sn['status']})"})

    if c["watchlist"] or c["trend"] == "deteriorating":
        pts, det = w["rating"]["watchlist_or_negative"], ("watchlist" if c["watchlist"] else "deteriorating")
    elif c["trend"] == "improving":
        pts, det = w["rating"]["improving"], "improving"
    else:
        pts, det = w["rating"]["stable"], "stable"
    parts.append({"factor": "Rating trend", "points": pts, "detail": f"{c['rating']} {det}"})

    ev = _events(cid)
    pts = min(w["events_cap"], sum(w["events"][e.split("|")[0]] for e in ev))
    parts.append({"factor": "Near-term events", "points": pts,
                  "detail": "; ".join(e.split("|")[1] for e in ev) or "none"})

    mag = abs(imp["ebitda_shock_pct"]) if imp["ebitda_shock_pct"] < 0 else 0.0
    pts = w["shock"]["ge10"] if mag >= 10 else w["shock"]["ge3"] if mag >= 3 else w["shock"]["ge1"] if mag >= 1 else w["shock"]["else"]
    parts.append({"factor": "Shock magnitude", "points": pts,
                  "detail": f"Adjusted EBITDA {imp['ebitda_shock_pct']:+.1f}%"})

    score = sum(p["points"] for p in parts)
    return {"score": score, "affected": score >= w["threshold"], "parts": parts}


def _health(cid: str) -> str:
    sev = severity(cid)
    imp = impact(cid)
    b = imp["binding"]
    if (b and b["band_after"] == "Breach") or CLIENTS[cid]["watchlist"]:
        return "at_risk"
    return "attention" if sev["affected"] else "healthy"


# --------------------------------------------------------------------------- #
# overview
# --------------------------------------------------------------------------- #

def _person(alias: str) -> dict:
    p = PEOPLE.get(alias, {})
    return {"alias": alias, "name": p.get("name", alias), "title": p.get("title", ""),
            "wall_status": p.get("wall_status", "")}


def overview() -> dict[str, Any]:
    tiers: dict[str, dict] = {}
    rels, reviews = [], []
    tot_exp = tot_commit = tot_rev = tot_rev_prior = 0.0
    health = {"healthy": 0, "attention": 0, "at_risk": 0}
    reviews_cert: dict[str, list[str]] = {}
    for cid, c in CLIENTS.items():
        x = EXTRA[cid]
        exp = sum(e["current"] for e in x["exposure_by_product"])
        commit = sum(fac["commitment"] for fac in c["facilities"])
        tot_exp += exp
        tot_commit += commit
        tot_rev += x["revenue_ytd_m"]
        tot_rev_prior += x["revenue_prior_ytd_m"]
        t = tiers.setdefault(c["tier"], {"tier": c["tier"], "exposure_m": 0.0, "count": 0})
        t["exposure_m"] += exp
        t["count"] += 1
        h = _health(cid)
        health[h] += 1
        sev = severity(cid)
        imp = impact(cid)
        rels.append({"id": cid, "name": c["name"], "sector": c["sector"], "tier": c["tier"],
                     "exposure_m": round(exp, 1), "revenue_ytd_m": x["revenue_ytd_m"],
                     "health": h, "score": sev["score"], "affected": sev["affected"],
                     "ebitda_shock_pct": imp["ebitda_shock_pct"],
                     "owner": _person(OWNERS[cid]["owner"])})

        for cov in c["covenants"]:
            if cov["next_test"]:
                cert = date.fromordinal(date.fromisoformat(cov["next_test"]).toordinal() + 45).isoformat()
                if (_days(cert) or 0) >= 0:
                    reviews_cert.setdefault(cert, []).append(c["name"])
                break
        d = _days(c["kyc_refresh_due"])
        if d is not None and 0 <= d <= 120:
            reviews.append({"date": c["kyc_refresh_due"], "client_id": cid, "client": c["name"],
                            "kind": "KYC refresh", "detail": f"Due in {d} days (CP-006)"})
        for lr in x["limit_reviews"]:
            d = _days(lr["next_review"])
            if d is not None and 0 <= d <= 120:
                reviews.append({"date": lr["next_review"], "client_id": cid, "client": c["name"],
                                "kind": "Limit review", "detail": lr["type"]})
    m = SKILL_DATA["meeting"]
    reviews.append({"date": m["date"], "client_id": HERO, "client": CLIENTS[HERO]["name"],
                    "kind": "Client meeting", "detail": f"{m['title']} - {m['time_local']}"})
    for cert, names in reviews_cert.items():
        reviews.append({"date": cert, "client_id": None,
                        "client": f"{len(names)} relationships" if len(names) > 1 else names[0],
                        "kind": "Q3 compliance certificates",
                        "detail": "Due ~45 days after the 30 Sep test date; Credit Risk recalculates (CP-005)"})
    reviews.sort(key=lambda r: r["date"])
    rels.sort(key=lambda r: -r["score"])
    affected = [r for r in rels if r["affected"]]

    return {
        "as_of": SCN["as_of"], "time_local": SCN["time_local"],
        "scenario": {"id": SCN["id"], "title": SCN["title"], "headline": SCN["headline"],
                     "alert": SCN["alert"], "pulse": SCN["pulse"]},
        "affected_count": len(affected),
        "kpis": {
            "exposure_m": round(tot_exp, 1), "committed_m": round(tot_commit, 1),
            "revenue_ytd_m": round(tot_rev, 2),
            "revenue_yoy_pct": round((tot_rev / tot_rev_prior - 1) * 100, 1) if tot_rev_prior else None,
            "relationships": len(CLIENTS), "health": health,
            "stressed_breaches": sum(1 for cid in CLIENTS
                                     if (impact(cid)["binding"] or {}).get("band_after") == "Breach"),
        },
        "tiers": sorted(({**t, "exposure_m": round(t["exposure_m"], 1)} for t in tiers.values()),
                        key=lambda t: t["tier"]),
        "relationships": rels,
        "reviews": reviews[:10],
        "sources": ["Fabric IQ", "Foundry IQ", "Work IQ", "Web IQ", "Market scenario (synthetic)"],
    }


# --------------------------------------------------------------------------- #
# triage + outreach
# --------------------------------------------------------------------------- #

def triage() -> dict[str, Any]:
    items = []
    for cid, c in CLIENTS.items():
        sev, imp = severity(cid), impact(cid)
        items.append({"id": cid, "name": c["name"], "tier": c["tier"], "sector": c["sector"],
                      "score": sev["score"], "affected": sev["affected"], "parts": sev["parts"],
                      "driver": imp["driver"], "ebitda_shock_pct": imp["ebitda_shock_pct"],
                      "binding": imp["binding"], "deal": imp.get("deal"), "pse": imp.get("pse")})
    items.sort(key=lambda i: -i["score"])
    return {"as_of": SCN["as_of"], "threshold": SCN["severity"]["threshold"],
            "method": SCN["severity"]["note"], "items": items}


def _outreach_message(cid: str) -> str:
    line = ("We'd like to share our committee's read on this morning's oil move and talk through "
            "what it means for your plans.")
    extra = ""
    if cid == HERO:
        # Public matters only: the restricted subject (sized hedging) is not mentioned.
        extra = " Happy to cover the facility discussion ahead of Thursday as well."
    return (f"Good morning - given the overnight move in crude, {line}{extra} "
            f"Would {SCN['clients'][cid]['slot']} work? - {_person(OWNERS[cid]['owner'])['name']}")


def outreach() -> dict[str, Any]:
    t = triage()
    plan = []
    for it in t["items"]:
        if not it["affected"]:
            continue
        s = SCN["clients"][it["id"]]
        owner = _person(OWNERS[it["id"]]["owner"])
        notes = []
        if it["id"] == HERO:
            notes.append("CR-2026-0184 restriction: public matters only; no hedging structure or solicitation (CP-008).")
        if it["binding"] and it["binding"]["band_after"] in ("Breach", "Tight"):
            notes.append("Loop in the sector Credit Risk director before any facility discussion (CP-005).")
        plan.append({"id": it["id"], "name": it["name"], "score": it["score"],
                     "contact_role": s["contact_role"], "channel": s["channel"],
                     "channel_detail": s["channel_detail"], "slot": s["slot"],
                     "caller": owner, "message": _outreach_message(it["id"]),
                     "guardrail_notes": notes})
    return {"as_of": SCN["as_of"], "items": plan,
            "channels": ["video", "phone", "chat", "email"]}


def approve_outreach(items: list[dict[str, Any]], upn: str) -> dict[str, Any]:
    out = []
    for i in items:
        cid = i.get("id")
        if cid not in SCN["clients"]:
            continue
        s = SCN["clients"][cid]
        changed = i.get("channel") not in (None, s["channel"]) or i.get("slot") not in (None, s["slot"])
        resp = s["response"]
        if changed:
            resp = {"status": "awaiting", "detail": f"Sent by {i.get('channel', s['channel'])} for {i.get('slot', s['slot'])}; awaiting reply."}
        out.append({"id": cid, "status": resp["status"], "detail": resp["detail"],
                    "channel": i.get("channel", s["channel"]), "slot": i.get("slot", s["slot"])})
        logger.info("AUDIT outreach approved by %s: %s via %s at %s",
                    upn, cid, i.get("channel", s["channel"]), i.get("slot", s["slot"]))
    return {"approved_by": upn, "items": out,
            "note": "Simulated delivery for the demo - no message leaves the tenant."}


# --------------------------------------------------------------------------- #
# client workspace
# --------------------------------------------------------------------------- #

def workspace(cid: str) -> dict[str, Any]:
    if cid not in CLIENTS:
        raise KeyError(cid)
    c, x, imp, sev = CLIENTS[cid], EXTRA[cid], impact(cid), severity(cid)
    sn = _sn(c)
    exp_total = sum(e["current"] for e in x["exposure_by_product"]) or 1.0
    by_line: dict[str, float] = {}
    for e in x["exposure_by_product"]:
        by_line[e["product"]] = by_line.get(e["product"], 0.0) + e["current"]
    maturities = sorted((fac for fac in c["facilities"] if fac["maturity"]), key=lambda f_: f_["maturity"])
    next_mat = maturities[0] if maturities else None
    posture = []
    if imp["binding"]:
        b = imp["binding"]
        posture.append({"label": "Covenant headroom (shocked)", "value": f"{b['headroom_after']:+.2f}x",
                        "sub": f"{b['covenant']}: {b['after']}x vs {b['threshold']}x", "tone": b["tone"]})
        posture.append({"label": "EBITDA cushion", "value": f"{b['cushion_after']}%",
                        "sub": f"was {b['cushion_before']}% before the shock",
                        "tone": "bad" if b["cushion_after"] < 5 else "warn" if b["cushion_after"] < 15 else "ok"})
    if sn:
        posture.append({"label": "Single-name limit", "value": f"{sn['utilisation_pct']}%",
                        "sub": f"CAD {sn['utilised']:,.1f}m of {sn['limit']:,.0f}m", "tone":
                        "bad" if sn["utilisation_pct"] >= 100 else "warn" if sn["utilisation_pct"] >= 85 else "ok"})
    posture.append({"label": "Internal rating", "value": c["rating"],
                    "sub": f"{c['trend']}{' - watchlist' if c['watchlist'] else ''}",
                    "tone": "bad" if c["watchlist"] else "warn" if c["trend"] == "deteriorating" else "ok"})
    kd = _days(c["kyc_refresh_due"])
    posture.append({"label": "KYC refresh", "value": f"{kd} days" if kd is not None else "-",
                    "sub": f"due {c['kyc_refresh_due']}", "tone": "warn" if kd is not None and kd <= 60 else "ok"})
    if next_mat:
        md = _days(next_mat["maturity"])
        posture.append({"label": "Next maturity", "value": next_mat["maturity"],
                        "sub": f"{next_mat['type']} - {md} days", "tone": "warn" if md is not None and md <= 365 else "ok"})
    if imp.get("pse"):
        p = imp["pse"]
        posture.append({"label": "Derivative PSE", "value": f"{p['pct_after']}%",
                        "sub": f"CAD {p['before']}m -> {p['after']}m of {p['limit']:.0f}m", "tone": "ok" if p["pct_after"] < 75 else "warn"})

    s = SCN["clients"][cid]
    return {
        "id": cid, "name": c["name"], "ticker": c["ticker"], "sector": c["sector"],
        "sub_sector": c["sub_sector"], "tier": c["tier"], "as_of": SCN["as_of"],
        "owner": _person(OWNERS[cid]["owner"]), "banker": _person(OWNERS[cid]["banker"]),
        "contact_role": s["contact_role"], "channel_detail": s["channel_detail"], "slot": s["slot"],
        "exposure_m": round(exp_total, 1), "committed_m": round(sum(f_["commitment"] for f_ in c["facilities"]), 1),
        "revenue_ytd_m": x["revenue_ytd_m"], "wallet_share_pct": x["wallet_share_pct"],
        "severity": sev, "impact": imp, "posture": posture,
        "allocation": [{"product": k, "exposure_m": round(v, 1), "pct": round(v / exp_total * 100, 1)}
                       for k, v in sorted(by_line.items(), key=lambda kv: -kv[1])],
        "engagement": x["engagement"], "trades": x["trades"],
        "restricted": cid == HERO,
        "prompts": prompts(cid),
    }


def prompts(cid: str) -> list[dict[str, str]]:
    c, imp = CLIENTS[cid], impact(cid)
    shock = (f"Scenario as at {SCN['as_of']} {SCN['time_local']} (synthetic): WTI -14.2% overnight on an "
             f"OPEC+ supply surprise, WCS-WTI differential -17.20, US energy HY OAS +62 bp, USD/CAD 1.392.")
    return [
        {"label": "Market context for this client",
         "prompt": f"{shock} What is the market and credit context for {c['name']} given this move, and what "
                   f"pressure points is the client likely to raise today? Treat the scenario as given."},
        {"label": "Exposure sensitivity + committee view",
         "prompt": f"{shock} Show {c['name']}'s covenant and limit sensitivity to a {imp['ebitda_shock_pct']:+.1f}% "
                   f"Adjusted EBITDA move, and the Credit & Markets Committee's current lending and hedging posture (CP-015)."},
        {"label": "My communication guardrails",
         "prompt": f"{shock} What are my communication guardrails for a call with {c['name']} today - what can I say, "
                   f"what needs a disclaimer, what needs approval, and what must I not say?"},
    ]


# --------------------------------------------------------------------------- #
# guardrails
# --------------------------------------------------------------------------- #

def guardrails(cid: str) -> dict[str, Any]:
    if cid not in CLIENTS:
        raise KeyError(cid)
    hv, c, imp = SCN["house_view"], CLIENTS[cid], impact(cid)
    can = list(hv["approved_framing"])
    disclaim = list(hv["needs_disclaimer"])
    approval = list(hv["needs_approval"])
    never = list(hv["never_say"])
    if cid == HERO:
        rx = SKILL_DATA["restriction"]
        never.insert(0, f"Any sized hedging structure or solicitation while {rx['id']} is live (expires {rx['expires']}) - CP-008.")
        can.append("General market conditions in the WCS differential - no structure, no notional.")
        approval.insert(0, f"Any revolver size or hold: the {imp['deal']['ask_m']:.0f}m ask now breaches at "
                           f"{imp['deal']['ask_leverage_after']}x on shocked EBITDA (CP-005 / CP-003).")
    if imp["binding"] and imp["binding"]["band_after"] in ("Breach", "Tight"):
        approval.append("Any covenant amendment, waiver or reset discussion - Credit Risk first (CP-005).")
    sources = [f"{FOUNDRY} - CP-014, CP-008, CP-003, CP-010", f"{FOUNDRY} - CP-015 house view", FABRIC]
    if cid == HERO:
        sources.append(f"{WORK} - CR-2026-0184")
    return {"id": cid, "name": c["name"], "can_say": can, "needs_disclaimer": disclaim,
            "needs_approval": approval, "never_say": never, "sources": sources,
            "house_view": hv["positions"]}


def approve_guardrails(cid: str, upn: str) -> dict[str, Any]:
    if cid not in CLIENTS:
        raise KeyError(cid)
    logger.info("AUDIT guardrails approved by %s for %s", upn, cid)
    return {"id": cid, "approved_by": upn, "approved": True}


# --------------------------------------------------------------------------- #
# call assist
# --------------------------------------------------------------------------- #

def call_script(cid: str) -> dict[str, Any]:
    if cid not in CLIENTS:
        raise KeyError(cid)
    c, imp, s = CLIENTS[cid], impact(cid), SCN["clients"][cid]
    b = imp["binding"]
    turns = [{
        "speaker": "Relationship manager", "kind": "open",
        "text": "Thanks for making the time. Our committee reads this morning's move as supply-driven, with a "
                "partial retracement over the next two quarters as the base case - so we'd focus on resilience "
                "rather than spot.",
        "sources": [f"{FOUNDRY} - CP-015"],
    }]
    if b:
        q = ("What does a move like this do to our "
             + ("leverage covenant?" if "Total Net Debt" in b["covenant"] else
                "interest coverage covenant?" if "Coverage" in b["covenant"] else
                "covenant headroom?"))
        a = (f"Illustratively, on your reported figures a {imp['ebitda_shock_pct']:+.1f}% change in Adjusted "
             f"EBITDA would take it from {b['before']}x to about {b['after']}x against the {b['threshold']}x "
             f"{'maximum' if b['test'] == 'maximum' else 'minimum'}. "
             + ("That would be a breach - worth getting ahead of with an early covenant conversation."
                if b["band_after"] == "Breach" else
                "That leaves limited room, so we'd suggest talking about covenant flexibility early."
                if b["band_after"] == "Tight" else
                "That still leaves reasonable room."))
        turns.append({"speaker": s["contact_role"], "kind": "question", "text": q})
        turns.append({"speaker": "Relationship manager", "kind": "answer", "text": a,
                      "disclaimer": "Illustrative, based on the client's reported figures.",
                      "sources": [f"{FABRIC} - covenants", "computed"]})
    if cid == HERO:
        d = imp["deal"]
        turns.append({"speaker": s["contact_role"], "kind": "question",
                      "text": "Does this change where you are on the revolver upsize for Thursday?"})
        turns.append({"speaker": "Relationship manager", "kind": "answer",
                      "text": "We're still keen to work on a larger facility and a longer tenor, and we think the "
                              "covenant conversation belongs alongside it. Any size is subject to our internal "
                              "approval - we'll come back with something firm rather than give you a number today.",
                      "internal": f"Shocked EBITDA cuts the largest clearing increase from {d['max_clearing_before_m']}m "
                                  f"to {d['max_clearing_after_m']}m; the resized +150m now runs {d['resized_150_leverage_after']}x "
                                  f"({'clears' if d['resized_150_clears'] else 'breaches'}). Re-cut the path to yes with Credit Risk.",
                      "sources": [f"{FOUNDRY} - CP-003, CP-005", "computed"]})
        turns.append({"speaker": s["contact_role"], "kind": "question",
                      "text": "With the differential this wide, should we be adding to our hedges now?"})
        turns.append({"speaker": "Relationship manager", "kind": "answer",
                      "text": "I can talk about what we're seeing in the differential market generally - liquidity and "
                              "where the curve has moved. We'll come back separately on anything specific.",
                      "internal": "CR-2026-0184 is live: general market conditions only; no sized structure, no solicitation (CP-008).",
                      "sources": [f"{WORK} - CR-2026-0184", f"{FOUNDRY} - CP-008"]})
    else:
        turns.append({"speaker": s["contact_role"], "kind": "question",
                      "text": "Would you give us more room on the facilities if we need it?"})
        turns.append({"speaker": "Relationship manager", "kind": "answer",
                      "text": "We're supportive and want to understand your plan first. Any change to facility size or "
                              "terms goes through our internal approval, so I'll come back with something firm rather "
                              "than a number today.",
                      "sources": [f"{FOUNDRY} - CP-003, CP-014"]})
        turns.append({"speaker": s["contact_role"], "kind": "question",
                      "text": "Should we be looking at our hedging after this move?"})
        turns.append({"speaker": "Relationship manager", "kind": "answer",
                      "text": "It's a good moment to review the programme. We can walk through it with you - any "
                              "recommendation follows our appropriateness process, so it would come as a separate proposal.",
                      "sources": [f"{FOUNDRY} - CP-010, CP-015"]})
    return {"id": cid, "name": c["name"], "contact_role": s["contact_role"],
            "channel_detail": s["channel_detail"], "slot": s["slot"], "turns": turns,
            "next_steps": next_steps(cid)}


def next_steps(cid: str) -> list[str]:
    c, imp = CLIENTS[cid], impact(cid)
    steps = ["Send a summary of today's call using the approved framing."]
    if cid == HERO:
        steps += ["Re-run the CP-005 pro-forma at shocked EBITDA with Credit Risk and re-cut the revolver path to yes.",
                  "Hold the hedging discussion until the Control Room clears CR-2026-0184.",
                  "Confirm the agenda for Thursday's CFO / Treasurer meeting."]
    else:
        b = imp["binding"]
        if b and b["band_after"] in ("Breach", "Tight"):
            steps.append("Model two covenant scenarios (base and -15% EBITDA) with Credit Risk.")
        steps.append("Offer a hedge-programme review, subject to our appropriateness process.")
        steps.append("Schedule a follow-up with the client after their next compliance certificate.")
    return steps


# --------------------------------------------------------------------------- #
# wrap-up (the "Cowork" step)
# --------------------------------------------------------------------------- #

# Anything a client-facing email must never carry (CP-008 / CP-014): restricted subject
# matter, information-barrier mechanics, internal process and the people on the wall.
_CLIENT_UNSAFE = ("control room", "cr-20", "restrict", "wall", "cross", "cp-",
                  "credit risk", "pro-forma", "proforma", "rating", "limit", "sasha", "ouellet",
                  "vance", "deleon")
# Only where a restriction is live: the restricted subject itself.
_RESTRICTED_UNSAFE = ("hedg", "differential", "overlay", "derivative")


def client_safe_steps(cid: str, agreed: list[str]) -> list[str]:
    """Only server-vetted next steps reach the client; free text never does."""
    vetted = set(next_steps(cid))
    banned = _CLIENT_UNSAFE + (_RESTRICTED_UNSAFE if cid == HERO else ())
    return [a for a in agreed if a in vetted and not any(w in a.lower() for w in banned)]


def wrapup(cid: str, agreed: list[str] | None, notes: str, upn: str) -> dict[str, Any]:
    if cid not in CLIENTS:
        raise KeyError(cid)
    ws, g = workspace(cid), guardrails(cid)
    imp = ws["impact"]
    agreed = [a.strip()[:300] for a in (agreed if agreed is not None else next_steps(cid)) if a.strip()][:12]
    client_steps = client_safe_steps(cid, agreed)
    b = imp["binding"]
    memo = {
        "title": f"Coverage memo - {ws['name']} - {SCN['as_of']}",
        "sections": [
            {"heading": "Context", "lines": [SCN["headline"], f"Client-specific driver: {imp['driver']}"]},
            {"heading": "Position (internal - CP-014)", "lines": [p["label"] + ": " + p["value"] + " - " + p["sub"] for p in ws["posture"]]},
            {"heading": "Shock impact", "lines": ([f"{b['covenant']}: {b['before']}x -> {b['after']}x vs {b['threshold']}x "
                                                   f"({b['band_after']}); EBITDA cushion {b['cushion_before']}% -> {b['cushion_after']}%."] if b else [])
                                                 + ([f"Revolver path to yes: largest clearing increase {imp['deal']['max_clearing_before_m']}m -> "
                                                     f"{imp['deal']['max_clearing_after_m']}m on shocked EBITDA."] if imp.get("deal") else [])},
            {"heading": "Call summary", "lines": [f"{ws['contact_role']} via {ws['channel_detail']} at {ws['slot']}. Approved framing used throughout."]
                                                 + ([f"Banker notes: {notes}"] if notes else [])},
            {"heading": "Agreed next steps", "lines": agreed},
            {"heading": "Guardrails applied", "lines": g["never_say"][:3]},
        ],
        "prepared_by": upn,
    }
    owner = ws["owner"]["name"]
    email = {
        "subject": f"Following up on today's call - {ws['name']}",
        "body": (f"Hello,\n\nThank you for making time this morning. As discussed, our committee reads the overnight "
                 f"move in crude as supply-driven, with a partial retracement over the next two quarters as the base case.\n\n"
                 f"Agreed next steps:\n" + ("\n".join(f"- {a}" for a in client_steps) or "- We will follow up shortly.") +
                 f"\n\nAny view on commodity prices or spreads is forward-looking and not a forecast or advice; any change "
                 f"to facility size or terms remains subject to our internal approval.\n\nBest regards,\n{owner}"),
        "to_note": "No recipient is set - add the client contact in Outlook before sending.",
    }
    logger.info("AUDIT wrap-up drafted by %s for %s", upn, cid)
    return {"id": cid, "name": ws["name"], "memo": memo, "email": email,
            "status": "DRAFT - prepared from governed data and the call. Review before sending.",
            "cowork_prompt": cowork_prompt(cid, agreed)}


def cowork_prompt(cid: str, agreed: list[str]) -> str:
    ws = workspace(cid)
    b = ws["impact"]["binding"]
    return (f"Using the CMCI Coverage Memo skill, prepare a coverage memo and a client follow-up email for "
            f"{ws['name']} following this morning's call with the {ws['contact_role']}. Context: {SCN['headline']} "
            + (f"{b['covenant']} moves from {b['before']}x to {b['after']}x against {b['threshold']}x on the shocked EBITDA. " if b else "")
            + "Agreed next steps: " + "; ".join(agreed) + ". Approved framing throughout; refer to our firm as "
              "\"the Bank\". Save both as drafts.")


def memo_docx(bundle: dict[str, Any]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Segoe UI"
    doc.styles["Normal"].font.size = Pt(10)
    r = doc.add_paragraph().add_run(bundle["status"])
    r.bold = True
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    m = bundle["memo"]
    doc.add_heading(m["title"], level=1)
    for s in m["sections"]:
        doc.add_heading(s["heading"], level=2)
        for line in s["lines"]:
            doc.add_paragraph(line, style="List Bullet")
    doc.add_heading("Client follow-up email (draft)", level=2)
    doc.add_paragraph(f"Subject: {bundle['email']['subject']}").runs[0].bold = True
    for para in bundle["email"]["body"].split("\n\n"):
        doc.add_paragraph(para)
    doc.add_paragraph().add_run(f"Prepared by {m['prepared_by']} - INTERNAL, confidential under CP-014.").italic = True
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
