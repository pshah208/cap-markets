"""Draft approval papers - the step after "path to yes".

Three drafts, built deterministically from the same governed figures the CP-005
skill uses, so every number in a paper reconciles to Fabric IQ and every rule to
a cited policy:

  * cc  - Credit Committee submission (CP-003, CP-005, CP-001, CP-006)
  * pm  - Portfolio Management sector offset / sell-down plan (CP-002)
  * cr  - Control Room clearance request (CP-008)

They are DRAFTS. Each carries an empty human sign-off block: the agent prepares a
decision, it never takes one. The firm is named "the Bank" throughout, and the
Control Room request deliberately names no individuals and no restricted subject
matter - CP-008 applies to the paper as much as to the agent.
"""

from __future__ import annotations

import io
from typing import Any

from .skill_data import SKILL_DATA
from .skills import STRESS_SHOCKS, proforma

FABRIC = "Fabric IQ"
FOUNDRY = "Foundry IQ"
WORK = "Work IQ"


def _cp003_authority(hold_m: float) -> str:
    if hold_m <= 75:
        return "Director, Credit Risk (sector) - 2 business days"
    if hold_m <= 150:
        return "Credit Committee (written resolution permitted) - 3 business days"
    if hold_m <= 250:
        return "Credit Committee (convened) - 5 business days"
    return "Credit Committee plus Chief Risk Officer - 5 business days, no expedited path"


def _band(pct: float) -> str:
    return "Breach" if pct >= 100 else "Amber" if pct >= 85 else "Green"


def build(incremental_m: float = 150.0, hold_m: float | None = None,
          final_hold_m: float | None = None) -> dict[str, Any]:
    d = SKILL_DATA
    rev = d["opportunities"]["revolver_upsize"]
    cur_fac = float(rev["current_commitment_total_m"])
    cur_hold = float(rev["bank_current_hold_m"])
    share = cur_hold / cur_fac
    inc = max(0.0, float(incremental_m))
    new_fac = cur_fac + inc
    hold = round(float(hold_m) if hold_m else new_fac * share, 1)
    final_hold = round(float(final_hold_m) if final_hold_m else cur_hold, 1)
    hold_inc = round(hold - cur_hold, 1)
    sell_down = round(max(0.0, hold - final_hold), 1)

    pf = proforma(inc)
    stress = [proforma(inc, s) for s in STRESS_SHOCKS]
    sn = next(l for l in d["limits"] if l["type"].startswith("Single-name"))
    pse = next(l for l in d["limits"] if "pre-settlement" in l["type"])
    sector = d["sector_limit"]
    m = d["metrics"]
    rx = d["restriction"]
    meeting = d["meeting"]
    hedge = d["opportunities"]["hedging_overlay"]
    as_of = d["client"]["as_of"]
    client = d["client"]["name"]

    sn_after = round(sn["utilised"] + max(0.0, hold_inc), 1)
    sn_pct = round(sn_after / sn["limit"] * 100, 1)
    sec_arranged = round(sector["utilised"] + max(0.0, hold_inc), 1)
    sec_final = round(sector["utilised"] + max(0.0, final_hold - cur_hold), 1)
    sec_arr_pct = round(sec_arranged / sector["limit"] * 100, 1)
    sec_fin_pct = round(sec_final / sector["limit"] * 100, 1)
    lev = next(c for c in d["covenants_all"] if "Total Net Debt" in c["covenant"])

    signoff = lambda *roles: [{"role": r, "decision": "", "name": "", "date": ""} for r in roles]

    cc = {
        "id": "cc",
        "title": f"Credit Committee submission - {client} revolver upsize",
        "addressee": "Credit Committee",
        "policies": ["CP-003", "CP-005", "CP-001", "CP-002", "CP-006"],
        "sections": [
            {"heading": "1. Request", "rows": [
                ["Obligor", f"{client} ({d['client']['ticker']})", FABRIC],
                ["Facility", f"Revolving credit facility, CAD {cur_fac:,.0f}m -> CAD {new_fac:,.1f}m (+{inc:,.1f}m) and one-year extension", FABRIC],
                ["The Bank's role", rev["bank_role"], FABRIC],
                ["The Bank's hold", f"CAD {cur_hold:,.1f}m -> CAD {hold:,.1f}m ({hold / new_fac * 100:.1f}% of facility)", FABRIC],
                ["Approval authority", _cp003_authority(hold), f"{FOUNDRY} - CP-003"],
                ["Decision needed by", rev["decision_needed_by"], FABRIC],
            ]},
            {"heading": "2. CP-005 pro-forma covenant test (fully drawn)", "rows": [
                ["Covenant", f"{lev['covenant']}, maximum {lev['threshold']}x", FABRIC],
                ["Current", f"{lev['actual']}x (prior quarter {lev['prior']}x) - {lev['headroom']} turns, {lev['status']}", FABRIC],
                ["Pro-forma", f"{pf['pro_forma_leverage']}x - {pf['headroom_turns']} turns, {pf['band']}", f"{FOUNDRY} - CP-005 (computed)"],
                ["Largest increase that clears", f"CAD {pf['max_clearing_m']:,.1f}m", "computed"],
                ["EBITDA cushion at this size", f"Adjusted EBITDA can fall {pf['ebitda_cushion_pct']}% before breach", "computed"],
            ]},
            {"heading": "3. EBITDA stress", "table": {
                "head": ["Adjusted EBITDA", "Leverage", "Headroom", "Band"],
                "rows": [[f"{s['ebitda_shock_pct']:+.0f}% (CAD {s['ebitda_m']:,.0f}m)",
                          f"{s['pro_forma_leverage']}x", f"{s['headroom_turns']:+.2f}", s["band"]]
                         for s in stress]},
             "source": "computed from Fabric IQ financials"},
            {"heading": "4. Limits after the increase", "rows": [
                ["Single-name", f"CAD {sn_after:,.1f}m of {sn['limit']:,.0f}m - {sn_pct}%, {_band(sn_pct)} (not the binding constraint)", f"{FABRIC} / {FOUNDRY} - CP-001"],
                ["Energy sector envelope", f"CAD {sec_arranged:,.1f}m of {sector['limit']:,.0f}m at arranged hold - {sec_arr_pct}%, {_band(sec_arr_pct)}", f"{FABRIC} / {FOUNDRY} - CP-002"],
                ["Derivative pre-settlement", f"{pse['utilisation_pct']}% used, CAD {pse['headroom']}m headroom - unaffected by this request", FABRIC],
            ]},
            {"heading": "5. Conditions", "lines": [
                [f"Sector is {sector['status']}: Portfolio Management sign-off and a documented distribution plan before commitment (see PM paper).", f"{FOUNDRY} - CP-002 / CP-003"],
                [f"KYC periodic refresh due {m['kyc_refresh_due']}: approval may proceed, but no documentation, execution or funding until the refresh completes.", f"{FOUNDRY} - CP-006"],
                ["Covenant relief or a step-down mechanic to be negotiated alongside the upsize, not after.", f"{WORK} - pro-forma sign-off"],
                ["No hold size or price indicated to the client before this approval.", f"{FOUNDRY} - CP-003"],
                [f"Restricted-list entry {rx['id']} in force until {rx['expires']}: pitch on public information only.", f"{WORK} / {FOUNDRY} - CP-008"],
            ]},
            {"heading": "6. Not decided here", "lines": [
                ["Allocation of sector capacity between this hold and the Global Markets hedging overlay is reserved to the Chief Risk Officer.", f"{FOUNDRY} - CP-002"],
            ]},
        ],
        "signoff": signoff("Sector Credit Risk director (pro-forma)", "Credit Committee chair"),
    }

    pm = {
        "id": "pm",
        "title": f"Portfolio Management - sector offset and distribution plan, {client}",
        "addressee": "Portfolio Management / Chief Risk Officer",
        "policies": ["CP-002", "CP-003"],
        "sections": [
            {"heading": "1. Sector position", "rows": [
                ["Energy sector envelope", f"CAD {sector['utilised']:,.0f}m of {sector['limit']:,.0f}m - {sector['utilisation_pct']}%, {sector['status']}", FABRIC],
                ["Owner", sector["owner"], f"{FABRIC} / {FOUNDRY} - CP-002"],
                ["Rule", "When the sector is Amber, new commitments need Portfolio Management sign-off and an identified offset; arranging above the final hold with a documented sell-down is the default.", f"{FOUNDRY} - CP-002"],
            ]},
            {"heading": "2. Proposed distribution", "rows": [
                ["Arranged position", f"CAD {hold:,.1f}m", "proposal"],
                ["Target final hold", f"CAD {final_hold:,.1f}m", "proposal"],
                ["Sell-down to place", f"CAD {sell_down:,.1f}m via Loan Syndications", "proposal"],
                ["Sector at arranged position", f"{sec_arr_pct}% ({_band(sec_arr_pct)})", "computed"],
                ["Sector after sell-down", f"{sec_fin_pct}% ({_band(sec_fin_pct)}) - {'sector-neutral' if final_hold <= cur_hold else 'net increase'}", "computed"],
            ]},
            {"heading": "3. To be completed by Loan Syndications", "lines": [
                ["Target investor list and indicative allocations.", "open"],
                ["Execution timeline relative to signing and the KYC refresh date.", "open"],
                ["Fallback if the sell-down is not placed (hold-to-maturity approval required).", f"{FOUNDRY} - CP-003"],
            ]},
            {"heading": "4. Correlation", "lines": [
                ["Several midstream obligors depend on the same egress corridor and are reported as a correlated cluster; adding to the largest is a concentration decision, not only a client decision.", f"{WORK} - sector position note"],
            ]},
        ],
        "signoff": signoff("Portfolio Management", "Chief Risk Officer (sector allocation)"),
    }

    cr = {
        "id": "cr",
        "title": f"Control Room clearance request - {rx['id']}",
        "addressee": "Control Room (Compliance)",
        "policies": ["CP-008", "CP-009", "CP-014"],
        "sections": [
            {"heading": "1. Restriction", "rows": [
                ["Reference", rx["id"], WORK],
                ["Raised", rx["raised_on"], WORK],
                ["Expected expiry", f"{rx['expires']} (an expectation, not a commitment)", WORK],
                ["Operational effect", rx["effect"], WORK],
            ]},
            {"heading": "2. Meeting", "rows": [
                ["Date / place", f"{meeting['date']} {meeting['time_local']} - {meeting['location']}", WORK],
                ["Cleared items", "1 relationship review; 2 revolver upsize and extension on public information; 5 sustainability-linked structures, indicative and unlabelled", WORK],
                ["Deferred / not cleared", "3 hedging overlay (no sized structure, no solicitation); 4 2027 refinancing (not to be raised by the Bank)", WORK],
            ]},
            {"heading": "3. Determinations requested", "lines": [
                [f"Re-shaped hedging structure (CAD {hedge['notional_m']:,.0f}m notional rebalanced toward the CAD/USD forward strip to fit CAD {pse['headroom']}m of pre-settlement headroom): may it be presented once the restriction lifts, or does it remain within the restricted subject matter?", f"{WORK} / {FOUNDRY} - CP-004"],
                ["Confirm the revolver discussion may reference a resized facility and a concurrent covenant conversation, on public information and the Bank's own credit analysis only.", f"{FOUNDRY} - CP-008"],
                ["Confirm the attendee list as cleared; any attendee with crossed status recuses if the conversation moves toward the restricted matter.", f"{FOUNDRY} - CP-008"],
            ]},
            {"heading": "4. Undertakings", "lines": [
                ["No solicitation, formal or informal, of any item not cleared.", f"{FOUNDRY} - CP-008"],
                ["No internal limits, ratings, sector capacity or wallet share disclosed to the client.", f"{FOUNDRY} - CP-014"],
                ["This request names no individuals and contains no restricted subject matter.", f"{FOUNDRY} - CP-008"],
            ]},
        ],
        "signoff": signoff("Requesting Managing Director", "Control Room officer"),
    }

    return {
        "as_of": as_of,
        "inputs": {"incremental_m": inc, "facility_m": new_fac, "hold_m": hold,
                   "final_hold_m": final_hold},
        "status": "DRAFT - prepared by the agent from governed data. Not approved. "
                  "Requires the human sign-offs below.",
        "confidential": "INTERNAL - confidential under CP-014. Not for disclosure to the client.",
        "papers": [cc, pm, cr],
    }


def to_docx(bundle: dict[str, Any], which: str = "all") -> bytes:
    """Render one paper (or all three) as a Word document."""
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from docx.shared import Pt, RGBColor

    papers = [p for p in bundle["papers"] if which in ("all", p["id"])]
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Segoe UI"
    style.font.size = Pt(10)

    for i, p in enumerate(papers):
        if i:
            doc.add_page_break()
        banner = doc.add_paragraph()
        run = banner.add_run(bundle["status"])
        run.bold = True
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        doc.add_heading(p["title"], level=1)
        meta = doc.add_paragraph()
        meta.add_run(f"To: {p['addressee']}    As at: {bundle['as_of']}    "
                     f"Policies: {', '.join(p['policies'])}").italic = True

        for s in p["sections"]:
            doc.add_heading(s["heading"], level=2)
            if "rows" in s:
                t = doc.add_table(rows=0, cols=3)
                t.style = "Light Grid Accent 1"
                for k, v, src in s["rows"]:
                    cells = t.add_row().cells
                    cells[0].text, cells[1].text, cells[2].text = k, v, src
                    cells[0].paragraphs[0].runs[0].bold = True
            if "table" in s:
                tb = s["table"]
                t = doc.add_table(rows=1, cols=len(tb["head"]))
                t.style = "Light Grid Accent 1"
                for c, h in zip(t.rows[0].cells, tb["head"]):
                    c.text = h
                for row in tb["rows"]:
                    for c, v in zip(t.add_row().cells, row):
                        c.text = str(v)
                if s.get("source"):
                    doc.add_paragraph().add_run(f"Source: {s['source']}").italic = True
            for text, src in s.get("lines", []):
                para = doc.add_paragraph(style="List Bullet")
                para.add_run(text)
                r = para.add_run(f"  [{src}]")
                r.italic = True
                r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

        doc.add_heading("Approval (human sign-off)", level=2)
        t = doc.add_table(rows=1, cols=4)
        t.style = "Light Grid Accent 1"
        for c, h in zip(t.rows[0].cells, ["Role", "Decision", "Name", "Date"]):
            c.text = h
        for s in p["signoff"]:
            t.add_row().cells[0].text = s["role"]
        foot = doc.add_paragraph()
        foot.add_run(bundle["confidential"]).italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
